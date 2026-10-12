"""Shift schedule (rules 4 and 5), loaded from .xlsx or .csv. Two layouts are understood.

1. The team roster (the real file, e.g. "data/Roster_L1_L2_2026_Schedule - L2 FS Team (Mid Oct - Dec).csv"):
   blocks titled like "6A-3P SHIFT (L2)", each with a header row
       Agent, Email, Role, Fixed Weekly Off, 9-Oct : Fri, 10-Oct : Sat, ...
   and one row per person with that day's shift code per date column: "6A-3P", "9P-6A" (ends next
   day), "Off", "EL", "Comp Off", ... People rotate between blocks, so the cell, not the block, is the
   shift. Cells we can't interpret (e.g. "swap(12th)") make that person-day unknown: no checks for it.
   Count rows and summary tables (no email) are ignored. Dates carry no year; it is inferred from the
   weekday in the header.

2. A simple long format (see config/shifts.example.csv), one row per engineer per shift:

    email,                 day,        start,  end
    pramod@emergent.sh,    Mon,        09:00,  18:00     <- weekly: repeats every Monday
    pramod@emergent.sh,    2026-10-12, 21:00,  06:00     <- one date; end before start = ends next day
    pramod@emergent.sh,    Sun,        off,              <- day off (or leave start empty)

`day` is a weekday name (Mon/Monday...) or a date (YYYY-MM-DD). Times are in SHIFTS["timezone"].
A dated row replaces the weekly rows for that engineer on that date.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

WEEKDAYS = {name: i for i, names in enumerate(
    [("mon", "monday"), ("tue", "tues", "tuesday"), ("wed", "wednesday"), ("thu", "thur", "thurs", "thursday"),
     ("fri", "friday"), ("sat", "saturday"), ("sun", "sunday")]) for name in names}


@dataclass(frozen=True)
class Shift:
    email: str
    start: datetime  # UTC
    end: datetime  # UTC

    def contains(self, t: datetime) -> bool:
        return self.start <= t < self.end


def _parse_time(v) -> time | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    s = str(v).strip().lower()
    if not s or s in {"off", "leave", "holiday", "-", "na", "n/a"}:
        return None
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I %p", "%I:%M%p", "%I%p"):
        try:
            return datetime.strptime(s.upper(), fmt).time()
        except ValueError:
            continue
    raise ValueError(f"unrecognised time {v!r}")


def _raw_rows(path: Path) -> list[list]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True, data_only=True).active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    with path.open(encoding="utf-8-sig", newline="") as f:
        return [list(r) for r in csv.reader(f)]


def _rows(raw: list[list]) -> list[dict]:
    header = [str(h or "").strip().lower() for h in raw[0]]
    return [dict(zip(header, r)) for r in raw[1:] if any(c not in (None, "") for c in r)]


# --- team roster grid ----------------------------------------------------------------------

SHIFT_CODE_RE = re.compile(r"^(\d{1,2})(?::(\d{2}))?\s*([AP])M?\s*-\s*(\d{1,2})(?::(\d{2}))?\s*([AP])M?$", re.I)
OFF_CODES = re.compile(r"^(off|wo|week ?off|el|cl|sl|pl|leave|holiday|comp ?off\b.*)$", re.I)
DATE_HEADER_RE = re.compile(r"^(\d{1,2})-([A-Za-z]{3})\w*\s*:\s*([A-Za-z]{3})", re.I)
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def parse_shift_code(code) -> tuple[time, time] | None | str:
    """"6A-3P" -> (06:00, 15:00); "Off"/"EL"/"Comp Off" -> None; anything else -> "unknown"."""
    s = str(code or "").strip()
    if not s or OFF_CODES.match(s):
        return None
    m = SHIFT_CODE_RE.match(s)
    if not m:
        return "unknown"

    def to_time(h, mins, ap):
        h = int(h) % 12 + (12 if ap.upper() == "P" else 0)
        return time(h, int(mins or 0))

    return to_time(m[1], m[2], m[3]), to_time(m[4], m[5], m[6])


def _header_date(label: str, around: datetime) -> date | None:
    m = DATE_HEADER_RE.match(str(label or "").strip())
    if not m or m[2].lower() not in MONTHS or m[3].lower()[:3] not in {k[:3] for k in WEEKDAYS}:
        return None
    day, month, weekday = int(m[1]), MONTHS[m[2].lower()], WEEKDAYS[m[3].lower()[:3]]
    candidates = []
    for year in (around.year - 1, around.year, around.year + 1):
        try:
            d = date(year, month, day)
        except ValueError:
            continue
        if d.weekday() == weekday:
            candidates.append(d)
    return min(candidates, key=lambda d: abs((d - around.date()).days)) if candidates else None


def is_roster_grid(raw: list[list]) -> bool:
    return any(
        len(r) > 4 and str(r[0] or "").strip().lower() == "agent" and str(r[1] or "").strip().lower() == "email"
        for r in raw
    )


def parse_roster_grid(raw: list[list], around: datetime):
    """-> (dated {(email, date): (start, end) | (None, None)}, unknown {(email, date)}, info {email: {...}})"""
    dated, unknown, info = {}, set(), {}
    section, columns = "", {}
    for r in raw:
        cells = [str(c).strip() if c is not None else "" for c in r]
        nonempty = [c for c in cells if c]
        if len(nonempty) == 1 and cells[0]:
            section = cells[0]  # block title, e.g. "6A-3P SHIFT (L2)" or "Doom-Loop Calls"
            continue
        if len(cells) > 4 and cells[0].lower() == "agent" and cells[1].lower() == "email":
            columns = {i: d for i, c in enumerate(cells) if (d := _header_date(c, around))}
            continue
        if len(cells) < 2 or "@" not in cells[1] or not columns:
            continue  # count rows, summary tables, blanks
        email = cells[1].lower()
        info[email] = {"name": cells[0], "role": cells[2] if len(cells) > 2 else "", "block": section}
        for i, d in columns.items():
            parsed = parse_shift_code(cells[i] if i < len(cells) else "")
            if parsed == "unknown":
                unknown.add((email, d))
            else:
                dated[(email, d)] = parsed if parsed else (None, None)
    return dated, unknown, info


class ShiftCalendar:
    def __init__(
        self,
        shifts: list[Shift],
        people: set[str] | None = None,
        unknown_days: set[tuple[str, date]] | None = None,
        info: dict[str, dict] | None = None,
        tz_name: str = "UTC",
    ):
        self.by_email: dict[str, list[Shift]] = {}
        for s in sorted(shifts, key=lambda s: s.start):
            self.by_email.setdefault(s.email, []).append(s)
        self.people = {p.lower() for p in (people or self.by_email)}
        self.unknown_days = unknown_days or set()
        self.info = info or {}
        self.tz = ZoneInfo(tz_name)

    @classmethod
    def load(cls, path: str | Path, tz_name: str, around: datetime, days: int = 14) -> "ShiftCalendar":
        """Expand the file into concrete shifts for `days` either side of `around`."""
        tz = ZoneInfo(tz_name)
        raw = _raw_rows(Path(path))
        weekly: dict[tuple[str, int], tuple[time | None, time | None]] = {}
        dated: dict[tuple[str, date], tuple[time | None, time | None]] = {}
        unknown: set[tuple[str, date]] = set()
        info: dict[str, dict] = {}
        if is_roster_grid(raw):
            dated, unknown, info = parse_roster_grid(raw, around.astimezone(tz))
        for i, r in enumerate(_rows(raw) if not info else [], start=2):
            email = str(r.get("email") or "").strip().lower()
            day = r.get("day") if r.get("day") is not None else r.get("date", r.get("weekday"))
            if not email or day in (None, ""):
                continue
            try:
                times = (_parse_time(r.get("start")), _parse_time(r.get("end")))
            except ValueError as e:
                raise ValueError(f"{path} row {i}: {e}") from None
            if isinstance(day, datetime):
                dated[(email, day.date())] = times
            elif isinstance(day, date):
                dated[(email, day)] = times
            elif str(day).strip().lower() in WEEKDAYS:
                weekly[(email, WEEKDAYS[str(day).strip().lower()])] = times
            else:
                dated[(email, date.fromisoformat(str(day).strip()))] = times

        emails = {e for e, _ in weekly} | {e for e, _ in dated} | {e for e, _ in unknown}
        local_today = around.astimezone(tz).date()
        shifts = []
        for email in emails:
            for offset in range(-days, days + 1):
                d = local_today + timedelta(days=offset)
                start_t, end_t = dated.get((email, d), weekly.get((email, d.weekday()), (None, None)))
                if start_t is None or end_t is None:
                    continue
                start = datetime.combine(d, start_t, tz)
                end = datetime.combine(d, end_t, tz)
                if end <= start:
                    end += timedelta(days=1)  # overnight shift
                shifts.append(Shift(email, start.astimezone(timezone.utc), end.astimezone(timezone.utc)))
        return cls(shifts, emails, unknown, info, tz_name)

    def has_schedule(self, email: str) -> bool:
        return email.lower() in self.people

    def is_certain(self, email: str, t: datetime) -> bool:
        """False when the roster cell for that day (or the day before, for overnight shifts) is unreadable,
        e.g. "swap(12th)": we don't know whether they were working, so don't judge them."""
        d = t.astimezone(self.tz).date()
        return (email.lower(), d) not in self.unknown_days and (email.lower(), d - timedelta(days=1)) not in self.unknown_days

    def shift_at(self, email: str, t: datetime) -> Shift | None:
        return next((s for s in self.by_email.get(email.lower(), []) if s.contains(t)), None)

    def on_shift(self, email: str, t: datetime) -> bool:
        return self.shift_at(email, t) is not None

    def last_end_before(self, email: str, t: datetime) -> datetime | None:
        ends = [s.end for s in self.by_email.get(email.lower(), []) if s.end <= t]
        return max(ends) if ends else None

    def boundaries(self, t0: datetime, t1: datetime) -> list[tuple[str, str, Shift]]:
        """Shift starts and ends with t0 < time <= t1, as (email, "start"|"end", shift)."""
        out = []
        for email, shifts in self.by_email.items():
            for s in shifts:
                if t0 < s.start <= t1:
                    out.append((email, "start", s))
                if t0 < s.end <= t1:
                    out.append((email, "end", s))
        return sorted(out, key=lambda b: b[2].start if b[1] == "start" else b[2].end)
